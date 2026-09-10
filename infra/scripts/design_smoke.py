"""Synthetic-only HTTP acceptance for durable style memory and themed generation."""

import json
import os
import time
from pathlib import Path
import httpx
from dotenv import dotenv_values
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]


def run():
    env = {**dotenv_values(ROOT / ".env"), **os.environ}
    output = ROOT / "test-results/design"
    output.mkdir(parents=True, exist_ok=True)
    reference = output / "synthetic-reference.pptx"
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("142D43")
    for x in (1, 5):
        box = slide.shapes.add_textbox(Inches(x), Inches(2.5), Inches(3), Inches(2))
        box.text = "SYNTHETIC REFERENCE ONLY"
        p = box.text_frame.paragraphs[0]
        p.font.size = Pt(32)
        p.font.color.rgb = RGBColor.from_string("69D2C5")
    ppt.save(reference)
    base = env.get("E2E_BASE_URL", "http://127.0.0.1:8080") + "/api/v1"
    with httpx.Client(base_url=base, timeout=30, trust_env=False) as client:

        def req(method, path, **kwargs):
            result = client.request(method, path, **kwargs)
            result.raise_for_status()
            return result.json()

        def wait(ident):
            deadline = time.monotonic() + 180
            while time.monotonic() < deadline:
                job = req("GET", "/jobs/" + ident)
                if job["state"] == "SUCCEEDED":
                    return job
                if job["state"] in {"FAILED", "CANCELLED", "WAITING_REVIEW"}:
                    raise RuntimeError(f"{job['step']}: {job['error_code']} {job['state']}")
                time.sleep(0.5)
            raise TimeoutError("Design acceptance job timed out")

        token = req("POST", "/auth/login", json={"username": "demo-user", "password": env["SEED_PASSWORD"]})
        client.headers["Authorization"] = "Bearer " + token["access_token"]
        project = req("POST", "/projects", json={"name": "Synthetic design acceptance"})
        with reference.open("rb") as source:
            uploaded = req(
                "POST",
                f"/projects/{project['id']}/design-references",
                files={"file": (reference.name, source, "application/octet-stream")},
            )
        learned = wait(uploaded["job"]["id"])
        theme_id = learned["result"]["template_id"]
        library = req("GET", "/design-library")
        theme = next(t for t in library["themes"] if t["id"] == theme_id)
        assert theme["config"]["accent"] == "69D2C5" and not library["model_fine_tuned"]
        assert "SYNTHETIC REFERENCE ONLY" not in json.dumps(theme["config"])
        req("PATCH", f"/projects/{project['id']}", json={"name": project["name"], "template_id": theme_id})
        with (ROOT / "tests/golden/fixtures/synthetic.hwp").open("rb") as source:
            document = req(
                "POST",
                f"/projects/{project['id']}/documents",
                files={"file": ("synthetic.hwp", source, "application/octet-stream")},
            )
        wait(document["job"]["id"])
        plan_id = wait(req("POST", f"/documents/{document['document_id']}/analyze")["id"])["result"][
            "plan_id"
        ]
        plan = req("GET", "/slide-plans/" + plan_id)
        assert plan["data"]["provenance"]["theme_id"] == theme_id
        generated = wait(req("POST", f"/slide-plans/{plan_id}/generate-ppt")["id"])["result"]
        assert generated["qa_status"] == "PASS"
        response = client.get(f"/artifacts/{generated['artifact_id']}/download")
        response.raise_for_status()
        path = output / "docker-learned-theme.pptx"
        path.write_bytes(response.content)
        deck = Presentation(path)
        assert str(deck.slides[0].background.fill.fore_color.rgb) == "142D43"
        assert all(
            "SYNTHETIC REFERENCE ONLY" not in s.text
            for slide in deck.slides
            for s in slide.shapes
            if s.has_text_frame
        )
        result = {
            "status": "PASS",
            "synthetic_only": True,
            "theme_id": theme_id,
            "project_id": project["id"],
            "plan_id": plan_id,
            "themes": len(library["themes"]),
            "source_text_copied": False,
            "learning": "style-feature-memory",
            "ppt_bytes": path.stat().st_size,
            "ppt_qa": "PASS",
            "sample": str(path),
        }
        (output / "http-result.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        print(json.dumps(result, indent=2))


if __name__ == "__main__":
    run()
