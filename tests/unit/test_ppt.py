import json
from pathlib import Path
from uuid import uuid4
from nps.contracts import ContentBlock, Evidence, PlannedSlide, SlidePlanContract
from nps.ppt import qa_ppt, render_ppt


def sample_plan():
    ref = Evidence(document_version_id=uuid4(), chunk_id=uuid4(), section="synthetic")
    return SlidePlanContract(
        document_versions=[ref.document_version_id],
        title="합성 검증 보고서",
        mock=True,
        provenance={"prompt_pack_version": "1.0"},
        slides=[
            PlannedSlide(
                order=1,
                title="합성 검증 현황",
                source_refs=[ref],
                content_blocks=[
                    ContentBlock(
                        type="table",
                        cells=[["항목", "건수"], ["검증 A", "12"], ["검증 B", "18"]],
                        source_refs=[ref],
                    )
                ],
            )
        ],
    ).model_dump(mode="json")


def test_editable_ppt_and_qa(tmp_path):
    plan = sample_plan()
    path = tmp_path / "sample.pptx"
    policy = json.loads(Path("templates/default.json").read_text(encoding="utf-8"))
    render_ppt(plan, path, policy)
    qa = qa_ppt(path, plan, policy)
    assert qa["status"] == "PASS", qa
    assert qa["editable_tables"] == 1 and qa["editable_text"] >= 2


def test_overflow_is_not_success(tmp_path):
    plan = sample_plan()
    plan["slides"][0]["content_blocks"] = [
        {"type": "text", "text": "한글 문장 " * 700, "source_refs": plan["slides"][0]["source_refs"]}
    ]
    path = tmp_path / "overflow.pptx"
    policy = json.loads(Path("templates/default.json").read_text(encoding="utf-8"))
    render_ppt(plan, path, policy)
    assert qa_ppt(path, plan, policy)["status"] == "FAIL"


def test_many_short_paragraphs_flow_without_content_loss(tmp_path):
    from pptx import Presentation
    from nps.rendering import preview_slide

    plan = sample_plan()
    text = "\n".join(f"합성 항목 {i}: {i}건" for i in range(35))
    plan["slides"][0]["content_blocks"] = [
        {"type": "text", "text": text, "source_refs": plan["slides"][0]["source_refs"]}
    ]
    path = tmp_path / "multicolumn.pptx"
    policy = json.loads(Path("templates/default.json").read_text(encoding="utf-8"))
    render_ppt(plan, path, policy)
    assert qa_ppt(path, plan, policy)["status"] == "PASS"
    boxes = [s.text for s in Presentation(path).slides[0].shapes if s.has_text_frame]
    assert "\n".join(boxes[1:-1]) == text
    preview_slide(plan["slides"][0], tmp_path / "preview.png", policy)
