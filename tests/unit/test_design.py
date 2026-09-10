import json
from pathlib import Path
from unittest.mock import patch
import fakeredis
import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from sqlalchemy import select
from nps.auth import issue_access
from nps.models import Template, Project
from nps.ppt import render_ppt, qa_ppt
from nps.themes import BASE, choose_theme
from nps.theme_extract import extract
from nps.pipeline import run_until_terminal
from nps.contracts import ContentBlock, Evidence, PlannedSlide, SlidePlanContract
from uuid import uuid4


def reference(path):
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("142D43")
    for x in (1, 5):
        box = slide.shapes.add_textbox(Inches(x), Inches(2.5), Inches(3), Inches(2))
        box.text = "SYNTHETIC PRIVATE REFERENCE TEXT"
        p = box.text_frame.paragraphs[0]
        p.font.name = "Malgun Gothic"
        p.font.size = Pt(32)
        p.font.color.rgb = RGBColor.from_string("69D2C5")
    ppt.save(path)


def test_style_memory_without_source_content(tmp_path):
    path = tmp_path / "synthetic.pptx"
    reference(path)
    style = extract(path)
    assert style["background"] == "142D43" and style["accent"] == "69D2C5"
    assert style["layout_arrangement"] == "columns"
    assert "SYNTHETIC PRIVATE REFERENCE TEXT" not in json.dumps(style)
    assert style["reference_slides"] == 1


def test_legacy_plan_keeps_matching_legacy_policy(db, users):
    from types import SimpleNamespace
    from nps.ppt import template_for

    project = db.scalar(select(Project).where(Project.owner_id == users["demo-user"].id))
    project.template_id = None
    template = template_for(db, SimpleNamespace(project_id=project.id, data={"provenance": {}}))
    assert template.config.get("design_engine") != "editorial-2"
    assert template.config["min_font_pt"] <= 17


def test_design_upload_pipeline_scope_and_reuse(client, db, users, tmp_path):
    user, other = users["demo-user"], users["other-user"]
    headers = {"Authorization": "Bearer " + issue_access(user)}
    foreign = {"Authorization": "Bearer " + issue_access(other)}
    project = db.scalar(select(Project).where(Project.owner_id == user.id))
    path = tmp_path / "reference.pptx"
    reference(path)
    url = f"/api/v1/projects/{project.id}/design-references"
    with patch("nps.jobs.broker", return_value=fakeredis.FakeRedis(decode_responses=True)):
        assert (
            client.post(url, headers=foreign, files={"file": ("ref.pptx", path.read_bytes())}).status_code
            == 403
        )
        ids = []
        for _ in range(2):
            response = client.post(url, headers=headers, files={"file": ("ref.pptx", path.read_bytes())})
            assert response.status_code == 202, response.text
            job = run_until_terminal(db, response.json()["job"]["id"])
            assert job.state == "SUCCEEDED", (job.step, job.error_code)
            ids.append(job.result["template_id"])
        assert ids[0] == ids[1]
        library = client.get("/api/v1/design-library", headers=headers).json()
        assert len(library["themes"]) == 23 and library["model_fine_tuned"] is False
        assert ids[0] not in {
            t["id"] for t in client.get("/api/v1/design-library", headers=foreign).json()["themes"]
        }
        assert choose_theme(db, project, {"title": "synthetic", "slides": []}).id == ids[0]
        foreign_project = db.scalar(select(Project).where(Project.owner_id == other.id))
        assert (
            client.patch(
                f"/api/v1/projects/{foreign_project.id}",
                headers=foreign,
                json={"name": "synthetic", "template_id": ids[0]},
            ).status_code
            == 404
        )
        # Validation failure must never create a learned template.
        bad = client.post(url, headers=headers, files={"file": ("bad.pptx", b"not a zip")})
        failed = run_until_terminal(db, bad.json()["job"]["id"])
        assert failed.state == "FAILED" and failed.error_code == "DOC_MIME_MISMATCH"
        assert len(db.scalars(select(Template).where(Template.org_id == user.org_id)).all()) == 1


def design_plan():
    ref = Evidence(document_version_id=uuid4(), chunk_id=uuid4(), section="합성 개발 자료")
    slides = []
    for layout in ("cover", "key_points", "comparison", "process", "table", "chart"):
        block = ContentBlock(
            text="목표: 문서 준비 시간을 줄입니다.\n핵심: 원문 근거를 연결해 정확성을 확인합니다.\n계획: 검증 결과를 반영해 단계적으로 개선합니다.",
            source_refs=[ref],
        )
        if layout in {"table", "chart"}:
            block = ContentBlock(
                type=layout,
                cells=[["항목", "건수"], ["합성 검증 A", "12"], ["합성 검증 B", "18"]],
                source_refs=[ref],
            )
        if layout == "cover":
            block.text = "합성 데이터로 검증한 발표자료 디자인"
        slides.append(
            PlannedSlide(
                order=len(slides) + 1,
                title="문서의 핵심을 발표자료로 연결합니다",
                layout_type=layout,
                content_blocks=[block],
                source_refs=[ref],
            )
        )
    return SlidePlanContract(
        title="합성 디자인 검증",
        document_versions=[ref.document_version_id],
        slides=slides,
        mock=True,
        provenance={"design_engine": "editorial-2"},
    ).model_dump(mode="json")


@pytest.mark.parametrize(
    "spec", json.loads(Path("templates/catalog.json").read_text(encoding="utf-8")), ids=lambda s: s["slug"]
)
def test_all_bundled_themes_have_editable_readable_layouts(spec, tmp_path):
    plan = design_plan()
    policy = {**BASE, **spec}
    path = tmp_path / "design.pptx"
    render_ppt(plan, path, policy)
    qa = qa_ppt(path, plan, policy)
    assert qa["status"] == "PASS", qa
    assert qa["editable_charts"] == 1 and qa["editable_tables"] == 1
    deck = Presentation(path)
    assert (
        len({tuple((s.left, s.top, s.width, s.height) for s in slide.shapes) for slide in deck.slides}) >= 4
    )


def test_modern_overflow_is_not_reported_as_success(tmp_path):
    plan = design_plan()
    plan["slides"][1]["content_blocks"][0]["text"] = "한글 내용 " * 900
    path = tmp_path / "overflow.pptx"
    render_ppt(plan, path, BASE)
    assert qa_ppt(path, plan, BASE)["status"] == "FAIL"


def test_original_image_is_preserved_and_missing_video_image_fails(tmp_path):
    from PIL import Image
    from nps.rendering import preview_slide
    from nps.errors import DomainError

    plan = design_plan()
    slide = plan["slides"][1]
    slide["layout_type"] = "image"
    slide["content_blocks"] = [dict(type="image", text="합성 원본", source_refs=slide["source_refs"])]
    path = tmp_path / "source.png"
    Image.new("RGB", (640, 320), "#0FAFBA").save(path)
    ppt = tmp_path / "image.pptx"
    render_ppt(plan, ppt, BASE, images={slide["slide_id"]: path})
    assert qa_ppt(ppt, plan, BASE)["status"] == "PASS"
    preview_slide(slide, tmp_path / "preview.png", BASE, path)
    with pytest.raises(DomainError, match="VIDEO_SOURCE_IMAGE_MISSING"):
        preview_slide(slide, tmp_path / "missing.png", BASE)


def test_editorial_selects_key_points_preserves_numbers_and_source():
    from nps.contracts import SemanticChunk, SourceLocation
    from nps.llm import MockLLMAdapter
    from nps.chunking import validate_evidence

    chunks = []
    for i in range(14):
        text = f"추진 목표 {i}: 합성 검증 결과 {120 + i}건을 확보했습니다.\n" + "\n".join(
            f"보조 설명 {n}: 개발 실험의 참고 문장입니다." for n in range(12)
        )
        chunks.append(
            SemanticChunk(
                document_id=uuid4(),
                version_id=uuid4(),
                section_id=f"s{i}",
                token_count=200,
                text=text,
                source_location=SourceLocation(section=f"s{i}"),
            )
        )
    output = MockLLMAdapter().generate(
        "slide_plan",
        {
            "document_title": "합성 개발 계획.hwp",
            "normalized": {"sections": []},
            "chunks": [c.model_dump(mode="json") for c in chunks],
        },
    )
    plan = SlidePlanContract.model_validate(output)
    assert len(plan.slides) == 11 and plan.provenance["omitted_chunk_count"] == 4
    assert (
        sum(len(b.text) for s in plan.slides for b in s.content_blocks) < sum(len(c.text) for c in chunks) / 2
    )
    assert validate_evidence(plan, chunks)["status"] == "PASS"
    assert all(s.content_blocks[0].text.count("\n") <= 2 for s in plan.slides)
    assert any("건을 확보했습니다" in (s.title + b.text) for s in plan.slides for b in s.content_blocks)


def test_vendored_theme_sources_match_pinned_checksums():
    import hashlib

    for theme in json.loads(Path("templates/catalog.json").read_text(encoding="utf-8")):
        if theme.get("source_sha256"):
            path = Path("templates/vendor/reveal") / (theme["slug"].removeprefix("reveal-") + ".css")
            assert hashlib.sha256(path.read_bytes()).hexdigest() == theme["source_sha256"]
            assert "Permission is hereby granted" in Path(theme["license_path"]).read_text(encoding="utf-8")
