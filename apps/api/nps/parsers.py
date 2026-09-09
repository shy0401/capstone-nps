"""Safe, bounded normalization adapters. No network fetching or embedded execution."""

import hashlib
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path
from defusedxml import ElementTree as ET
from nps.config import settings
from nps.contracts import Node, NormalizedDocument, Section, SourceLocation
from nps.errors import DomainError
from nps.hwp import HWPParser
from nps.storage import storage
from nps.upload import archive_guard


def local_name(element):
    return element.tag.rsplit("}", 1)[-1]


def text_nodes(element):
    return "".join(x.text or "" for x in element.iter() if local_name(x) in {"t", "text"})


def save_image(data):
    from PIL import Image
    import io

    try:
        with Image.open(io.BytesIO(data)) as image:
            if image.width * image.height > 25_000_000:
                raise DomainError("DOC_RESOURCE_LIMIT")
            image.verify()
        key, path = storage.allocate("visuals", ".png")
        with Image.open(io.BytesIO(data)) as image:
            image.convert("RGB").save(path, format="PNG")
        return key, hashlib.sha256(data).hexdigest()
    except DomainError:
        raise
    except Exception:
        return None, hashlib.sha256(data).hexdigest()


class PDFParser:
    version = "pdf-1.0"

    def parse(self, path):
        import pdfplumber
        from pypdf import PdfReader

        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise DomainError("DOC_ENCRYPTED")
        if len(reader.pages) > settings().max_pages:
            raise DomainError("DOC_RESOURCE_LIMIT")
        sections, warnings = [], []
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                section_id = f"page-{i}"
                loc = SourceLocation(page=i, section=section_id)
                tables = page.find_tables()
                nodes = []

                # Table regions do not also appear as duplicate prose.
                def outside_tables(obj):
                    x, y = obj.get("x0", 0), obj.get("top", 0)
                    return not any(
                        t.bbox[0] <= x <= t.bbox[2] and t.bbox[1] <= y <= t.bbox[3] for t in tables
                    )

                prose = page.filter(outside_tables).extract_text() or ""
                for j, line in enumerate(prose.splitlines()):
                    nodes.append(
                        Node(
                            type="heading" if j == 0 else "paragraph",
                            text=line,
                            source_location=loc.model_copy(update={"paragraph": j}),
                        )
                    )
                for j, table in enumerate(tables):
                    cells = [[str(c or "") for c in row] for row in table.extract()]
                    nodes.append(
                        Node(
                            type="table",
                            cells=cells,
                            header=True,
                            text="\n".join(" | ".join(r) for r in cells),
                            source_location=loc.model_copy(update={"table": j}),
                        )
                    )
                for image in reader.pages[i - 1].images:
                    key, digest = save_image(image.data)
                    nodes.append(
                        Node(
                            type="image",
                            text=image.name,
                            source_location=loc,
                            image_key=key,
                            image_sha256=digest,
                        )
                    )
                    if key is None:
                        warnings.append("IMAGE_FORMAT_NOT_RENDERABLE")
                sections.append(
                    Section(
                        section_id=section_id,
                        heading=prose.splitlines()[0][:80] if prose else section_id,
                        nodes=nodes,
                    )
                )
        if not any(s.nodes for s in sections):
            raise DomainError("DOC_OCR_REQUIRED")
        return NormalizedDocument(
            metadata={"format": "pdf", "pages": len(sections)},
            sections=sections,
            parser_version=self.version,
            warnings=warnings,
        )


class DOCXParser:
    version = "docx-1.0"

    def parse(self, path):
        from docx import Document as Docx
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        archive_guard(path)
        doc = Docx(path)
        sections = [Section(section_id="section-1", heading="문서", nodes=[])]
        paragraph_i = table_i = 0
        for child in doc.element.body:
            tag = local_name(child)
            section = sections[-1]
            if tag == "p":
                paragraph = Paragraph(child, doc)
                is_heading = bool(paragraph.style and paragraph.style.name.startswith("Heading"))
                if is_heading:
                    section = Section(
                        section_id=f"section-{len(sections) + 1}",
                        heading=paragraph.text[:80] or "제목",
                        nodes=[],
                    )
                    sections.append(section)
                loc = SourceLocation(section=section.section_id, paragraph=paragraph_i)
                paragraph_i += 1
                if paragraph_i > settings().max_cells:
                    raise DomainError("DOC_RESOURCE_LIMIT")
                if paragraph.text:
                    section.nodes.append(
                        Node(
                            type="heading" if is_heading else "paragraph",
                            text=paragraph.text,
                            source_location=loc,
                        )
                    )
                for element in child.iter():
                    if local_name(element) == "blip":
                        rel_id = element.get(
                            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                        )
                        if rel_id and rel_id in doc.part.related_parts:
                            key, digest = save_image(doc.part.related_parts[rel_id].blob)
                            section.nodes.append(
                                Node(type="image", source_location=loc, image_key=key, image_sha256=digest)
                            )
            elif tag == "tbl":
                table = Table(child, doc)
                cells = [[c.text for c in row.cells] for row in table.rows]
                if sum(len(r) for r in cells) > settings().max_cells:
                    raise DomainError("DOC_RESOURCE_LIMIT")
                merges = []
                for r, row in enumerate(table.rows):
                    for c, cell in enumerate(row.cells):
                        props = cell._tc.tcPr
                        if props is not None and (props.gridSpan is not None or props.vMerge is not None):
                            merges.append(
                                {
                                    "row": r,
                                    "col": c,
                                    "colspan": cell._tc.grid_span,
                                    "vertical": props.vMerge.val if props.vMerge is not None else None,
                                }
                            )
                section.nodes.append(
                    Node(
                        type="table",
                        cells=cells,
                        header=True,
                        merges=merges,
                        text="\n".join(" | ".join(row) for row in cells),
                        source_location=SourceLocation(section=section.section_id, table=table_i),
                    )
                )
                table_i += 1
        return NormalizedDocument(
            metadata={"format": "docx"},
            sections=[s for s in sections if s.nodes],
            parser_version=self.version,
        )


class XLSXParser:
    version = "xlsx-1.0"

    def parse(self, path):
        from openpyxl import load_workbook

        archive_guard(path)
        source = path.open("rb")
        wb = load_workbook(source, data_only=True, keep_links=False)
        try:
            if len(wb.worksheets) > settings().max_sheets:
                raise DomainError("DOC_RESOURCE_LIMIT")
            sections = []
            for i, sheet in enumerate(wb.worksheets):
                if sheet.max_row * sheet.max_column > settings().max_cells:
                    raise DomainError("DOC_RESOURCE_LIMIT")
                cells = [
                    [str(c if c is not None else "") for c in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                sid = f"sheet-{i + 1}"
                loc = SourceLocation(section=sid, sheet=sheet.title, table=0)
                nodes = [
                    Node(type="heading", text=sheet.title, source_location=loc),
                    Node(
                        type="table",
                        cells=cells,
                        header=True,
                        text="\n".join(" | ".join(row) for row in cells),
                        source_location=loc,
                        merges=[{"range": str(r)} for r in sheet.merged_cells.ranges],
                    ),
                ]
                for image in sheet._images:
                    key, digest = save_image(image._data())
                    nodes.append(Node(type="image", image_key=key, image_sha256=digest, source_location=loc))
                sections.append(Section(section_id=sid, heading=sheet.title, nodes=nodes))
            return NormalizedDocument(
                metadata={"format": "xlsx", "formula_policy": "cached values only"},
                sections=sections,
                parser_version=self.version,
                warnings=["FORMULA_VALUES_REQUIRE_SAVED_CACHE"],
            )
        finally:
            wb.close()
            source.close()


class HWPXParser:
    version = "hwpx-1.0"

    def parse(self, path):
        names = archive_guard(path)
        section_names = sorted(n for n in names if n.startswith("Contents/section") and n.endswith(".xml"))
        if len(section_names) > settings().max_pages:
            raise DomainError("DOC_RESOURCE_LIMIT")
        sections = []
        with zipfile.ZipFile(path) as z:
            for i, name in enumerate(section_names):
                root = ET.fromstring(z.read(name))
                sid = f"section-{i + 1}"
                nodes, seen = [], set()
                for element in root.iter():
                    if id(element) in seen:
                        continue
                    tag = local_name(element)
                    loc = SourceLocation(section=sid, paragraph=len(nodes))
                    if tag == "tbl":
                        rows, merges = [], []
                        for ri, tr in enumerate(x for x in element if local_name(x) == "tr"):
                            row = []
                            for ci, tc in enumerate(x for x in tr if local_name(x) == "tc"):
                                row.append(text_nodes(tc))
                                spans = next((x for x in tc if local_name(x) == "cellSpan"), None)
                                if spans is not None:
                                    merges.append({"row": ri, "col": ci, **spans.attrib})
                            rows.append(row)
                        nodes.append(
                            Node(
                                type="table",
                                cells=rows,
                                header=True,
                                merges=merges,
                                text="\n".join(" | ".join(r) for r in rows),
                                source_location=loc,
                            )
                        )
                        seen.update(id(x) for x in element.iter())
                    elif tag == "p" and not any(local_name(x) == "tbl" for x in element.iter()):
                        value = text_nodes(element)
                        if value:
                            nodes.append(Node(type="paragraph", text=value, source_location=loc))
                    if len(nodes) > settings().max_cells:
                        raise DomainError("DOC_RESOURCE_LIMIT")
                # Binary image references remain tied to their section. No href is fetched.
                for img in (x for x in root.iter() if local_name(x) in {"img", "image"}):
                    binary_id = img.get("binaryItemIDRef", "")
                    candidates = [n for n in names if n.startswith("BinData/") and Path(n).stem == binary_id]
                    if candidates:
                        key, digest = save_image(z.read(candidates[0]))
                        nodes.append(
                            Node(
                                type="image",
                                image_key=key,
                                image_sha256=digest,
                                source_location=SourceLocation(section=sid),
                            )
                        )
                sections.append(
                    Section(
                        section_id=sid, heading=next((n.text[:80] for n in nodes if n.text), sid), nodes=nodes
                    )
                )
        return NormalizedDocument(metadata={"format": "hwpx"}, sections=sections, parser_version=self.version)


class ParserRegistry:
    adapters = {
        ".pdf": PDFParser(),
        ".docx": DOCXParser(),
        ".xlsx": XLSXParser(),
        ".hwpx": HWPXParser(),
        ".hwp": HWPParser(),
    }

    def parse(self, path, ext):
        if ext not in self.adapters:
            raise DomainError("DOC_UNSUPPORTED")
        try:
            result = self.adapters[ext].parse(path)
            if not result.sections or not any(s.nodes for s in result.sections):
                raise DomainError("DOC_EMPTY")
            return result
        except DomainError:
            raise
        except Exception:
            raise DomainError("DOC_PARSE_FAILED", 422) from None


def safe_parse(path, ext):
    """Subprocess timeout + Linux address-space limit. API never parses office files."""
    key, output = storage.allocate("parsed", ".json")
    env = os.environ.copy()
    env["PYTHONPATH"] = str(Path(__file__).resolve().parents[1])
    env["STORAGE_ROOT"] = str(settings().storage_root.resolve())
    try:
        result = subprocess.run(
            [sys.executable, "-m", "nps.parse_runner", str(path.resolve()), ext, str(output.resolve())],
            env=env,
            capture_output=True,
            timeout=settings().parse_timeout + 5,
            check=False,
        )
        if not output.exists():
            raise DomainError("DOC_PARSE_FAILED")
        data = json.loads(output.read_text(encoding="utf-8"))
        if result.returncode or "error" in data:
            raise DomainError(data.get("error", "DOC_PARSE_FAILED"), 422)
        return NormalizedDocument.model_validate(data)
    except subprocess.TimeoutExpired:
        raise DomainError("DOC_PARSE_TIMEOUT", 422) from None
    finally:
        output.unlink(missing_ok=True)
