import json
import os
import zipfile
import hashlib
from pathlib import Path
from nps.db import utcnow
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from sqlalchemy import func, select
from nps.auth import audit
from nps.errors import DomainError
from nps.models import Artifact, ArtifactVersion, Project, User, VisualAsset
from nps.rendering import text_fits, text_layout
from nps.storage import sha256_file, storage


def template_for(db, plan):
    from nps.themes import plan_theme

    selected = plan_theme(db, plan)
    if selected:
        from types import SimpleNamespace

        return SimpleNamespace(
            id=selected.id,
            version=selected.version,
            official_flag=False,
            storage_key=None,
            config=plan.data["provenance"]["theme_policy"],
        )
    project = db.get(Project, plan.project_id)
    from nps.themes import scoped_template, visible_templates

    template = scoped_template(db, project.template_id, project.org_id) if project.template_id else None
    # Pre-v2 plans have no frozen style: keep their legacy renderer/policy pair.
    if not template or template.config.get("design_engine") == "editorial-2":
        template = next(
            (
                t
                for t in visible_templates(db, project.org_id)
                if t.config.get("design_engine") != "editorial-2"
            ),
            None,
        )
    if not template:
        raise DomainError("TEMPLATE_NOT_FOUND", 404)
    return template


def add_text(slide, text, x, y, width, height, size, policy):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.06)
    frame.margin_top = frame.margin_bottom = Inches(0.04)
    frame.text = text
    for paragraph in frame.paragraphs:
        paragraph.font.name = policy["font"]
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = RGBColor.from_string(policy["foreground"])
        paragraph.line_spacing = 1.2
    return box


def render_ppt(plan, output, policy, template_path=None, images=None):
    if plan.get("provenance", {}).get("design_engine") == "editorial-2":
        from nps.presentation_design import render

        return render(plan, output, policy, images)
    deck = Presentation(template_path) if template_path else Presentation()
    # External template masters/layouts are preserved, example slides are removed.
    for slide_id in list(deck.slides._sldIdLst):
        deck.part.drop_rel(slide_id.rId)
        deck.slides._sldIdLst.remove(slide_id)
    deck.slide_width, deck.slide_height = Inches(policy["width_inches"]), Inches(policy["height_inches"])
    for item in plan["slides"]:
        slide = deck.slides.add_slide(deck.slide_layouts[6 if len(deck.slide_layouts) > 6 else 0])
        for shape in list(slide.shapes):
            if shape.is_placeholder:
                shape._element.getparent().remove(shape._element)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor.from_string(policy["background"])
        add_text(slide, item["title"], 0.65, 0.55, 12, 0.8, policy["title_pt"], policy)
        y = 1.65
        available = 4.7 / len(item["content_blocks"])
        decorative = (
            (images or {}).get(item["slide_id"])
            if all(b["type"] == "text" for b in item["content_blocks"])
            else None
        )
        for block in item["content_blocks"]:
            if block["type"] in {"table", "chart"}:
                cells = block["cells"]
                if block["type"] == "chart":
                    try:
                        data = CategoryChartData()
                        data.categories = [r[0] for r in cells[1:]]
                        for col in range(1, len(cells[0])):
                            data.add_series(cells[0][col], [float(r[col]) for r in cells[1:]])
                        slide.shapes.add_chart(
                            XL_CHART_TYPE.COLUMN_CLUSTERED,
                            Inches(0.65),
                            Inches(y),
                            Inches(12),
                            Inches(available - 0.1),
                            data,
                        )
                    except ValueError:
                        raise DomainError("CHART_NON_NUMERIC", 422) from None
                else:
                    table = slide.shapes.add_table(
                        len(cells),
                        len(cells[0]),
                        Inches(0.65),
                        Inches(y),
                        Inches(12),
                        Inches(min(available - 0.1, len(cells) * 0.55)),
                    ).table
                    for ri, row in enumerate(cells):
                        for ci, value in enumerate(row):
                            cell = table.cell(ri, ci)
                            cell.text = value
                            cell.margin_left = cell.margin_right = Inches(0.07)
                            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = RGBColor.from_string(
                                policy["accent"] if ri == 0 else "EDF3F3"
                            )
                            for paragraph in cell.text_frame.paragraphs:
                                paragraph.font.name = policy["font"]
                                paragraph.font.size = Pt(policy["min_font_pt"])
                                paragraph.font.color.rgb = RGBColor.from_string(
                                    "FFFFFF" if ri == 0 else policy["foreground"]
                                )
            elif block["type"] == "image" and (images or {}).get(item["slide_id"]):
                from PIL import Image

                image_path = images[item["slide_id"]]
                with Image.open(image_path) as im:
                    aspect = im.width / im.height
                width, height = min(12, (available - 0.1) * aspect), min(available - 0.1, 12 / aspect)
                slide.shapes.add_picture(
                    str(image_path), Inches(0.65), Inches(y), width=Inches(width), height=Inches(height)
                )
            else:
                size, width, parts = text_layout(
                    block.get("text", ""),
                    7.3 if decorative else 12,
                    available - 0.1,
                    policy["body_pt"],
                    policy["min_font_pt"],
                )
                for index, part in enumerate(parts):
                    add_text(
                        slide, part, 0.65 + index * (width + 0.3), y, width, available - 0.1, size, policy
                    )
            y += available
        if decorative:
            from PIL import Image

            with Image.open(decorative) as image:
                ratio = image.width / image.height
            width, height = min(4.2, 4.2 * ratio), min(4.2, 4.2 / ratio)
            slide.shapes.add_picture(
                str(decorative), Inches(8.4), Inches(1.85), width=Inches(width), height=Inches(height)
            )
        # Generated visuals do not replace editable content.
        slide.notes_slide.notes_text_frame.text = json.dumps(
            {
                "slide_id": item["slide_id"],
                "source_refs": item["source_refs"],
                "mock": plan["mock"],
                "provenance": plan["provenance"],
            },
            ensure_ascii=False,
        )
        add_text(
            slide,
            "내부 개발용 · 검토 필요" + (" · MOCK" if plan["mock"] else ""),
            0.65,
            6.55,
            12,
            0.4,
            17,
            policy,
        )
    deck.core_properties.title = plan["title"]
    deck.core_properties.subject = "Internal prototype; official template TBD-NPS-OUT-001"
    deck.save(output)


def qa_ppt(path, plan, policy):
    issues = []
    with zipfile.ZipFile(path) as z:
        if z.testzip() is not None or "[Content_Types].xml" not in z.namelist():
            issues.append("PPT-STRUCT")
    deck = Presentation(path)
    if len(deck.slides) != len(plan["slides"]):
        issues.append("PPT-SLIDE-COUNT")
    editable_text = editable_tables = editable_charts = 0
    for index, slide in enumerate(deck.slides):
        expected_images = sum(b["type"] == "image" for b in plan["slides"][index]["content_blocks"])
        if sum(shape.shape_type == 13 for shape in slide.shapes) < expected_images:
            issues.append(f"PPT-IMAGE-MISSING:{index + 1}")
        for shape in slide.shapes:
            margin = Inches(policy["safe_margin_inches"])
            if (
                shape.left < margin
                or shape.top < margin
                or shape.left + shape.width > deck.slide_width - margin
                or shape.top + shape.height > deck.slide_height - margin
            ):
                issues.append(f"PPT-BOUND:{index + 1}")
            if shape.has_text_frame and not shape.name.startswith("decor:"):
                editable_text += 1
                for paragraph in shape.text_frame.paragraphs:
                    size = paragraph.font.size.pt if paragraph.font.size else policy["body_pt"]
                    if size < (11 if shape.name.startswith("meta:") else policy["min_font_pt"]):
                        issues.append(f"PPT-FONT:{index + 1}")
                size = max(
                    (
                        p.font.size.pt if p.font.size else policy["body_pt"]
                        for p in shape.text_frame.paragraphs
                    ),
                    default=22,
                )
                if not text_fits(shape.text, size, shape.width.inches - 0.12, shape.height.inches - 0.08):
                    issues.append(f"PPT-TEXT:{index + 1}")
            if shape.has_table:
                editable_tables += 1
                for ri, row in enumerate(shape.table.rows):
                    for ci, cell in enumerate(row.cells):
                        if not text_fits(
                            cell.text,
                            policy["min_font_pt"],
                            shape.table.columns[ci].width.inches - 0.14,
                            row.height.inches - 0.08,
                        ):
                            issues.append(f"PPT-TABLE-TEXT:{index + 1}:{ri}:{ci}")
            if shape.has_chart:
                editable_charts += 1
            if shape.shape_type == 13:
                source_ratio = shape.image.size[0] / shape.image.size[1]
                if abs((shape.width / shape.height) / source_ratio - 1) > 0.01:
                    issues.append(f"PPT-IMAGE-ASPECT:{index + 1}")
        if not all(block["source_refs"] for block in plan["slides"][index]["content_blocks"]):
            issues.append("PPT-EVID")
        if not slide.notes_slide.notes_text_frame.text:
            issues.append("PPT-PROV")
    return {
        "status": "FAIL" if issues else "PASS",
        "issues": sorted(set(issues)),
        "editable_text": editable_text,
        "editable_tables": editable_tables,
        "editable_charts": editable_charts,
        "slide_count": len(deck.slides),
        "structure_check": "ZIP CRC + python-pptx reopen",
        "office_application_smoke": "NOT_RUN",
        "text_check": "font-metric conservative wrapping; renderer review still required",
        "official_template": policy["official_flag"],
        "schema_version": "1.0",
    }


def render_step(db, job, plan, cancelled):
    template = template_for(db, plan)
    key, path = storage.allocate("artifacts", ".pptx")
    images = {}
    for asset in db.scalars(select(VisualAsset).where(VisualAsset.plan_id == plan.id)):
        images[asset.slide_id] = storage.path(asset.storage_key)
    cancelled()
    render_ppt(
        plan.data,
        path,
        template.config,
        storage.path(template.storage_key) if template.storage_key else None,
        images,
    )
    return {"storage_key": key, "template_id": template.id}


def persist_artifact(db, job, plan, kind, key, qa, extra):
    template = template_for(db, plan)
    artifact = db.get(Artifact, job.payload.get("artifact_id")) if job.payload.get("artifact_id") else None
    if not artifact:
        artifact = Artifact(project_id=plan.project_id, plan_id=plan.id, type=kind)
        db.add(artifact)
        db.flush()
    db.refresh(artifact, with_for_update=True)
    number = (
        db.scalar(select(func.max(ArtifactVersion.version)).where(ArtifactVersion.artifact_id == artifact.id))
        or 0
    ) + 1
    path = storage.path(key)
    qa_key = storage.write_json("qa", qa)
    provenance = {
        **plan.data["provenance"],
        "app_version": "0.2.0",
        "plan_id": plan.id,
        "plan_version": plan.version,
        "review_mode": job.payload.get("review_mode", "strict"),
        "template_id": template.id,
        "template_version": template.version,
        "official_template": template.official_flag,
        "document_version_ids": plan.data["document_versions"],
        "generated_by_job_id": job.id,
        "generated_at": utcnow().isoformat(),
        "git_sha": os.environ.get("APP_GIT_SHA", "UNKNOWN"),
        "workflow_version": "wf-image-1.0",
        "workflow_hash": hashlib.sha256(Path("workflows/comfy/wf-image-v1.json").read_bytes()).hexdigest(),
        "visual_assets": [
            a.provenance for a in db.scalars(select(VisualAsset).where(VisualAsset.plan_id == plan.id))
        ],
        "slides": plan.data["slides"],
        "mock": plan.data["mock"],
        "qa_key": qa_key,
        **extra,
    }
    version = ArtifactVersion(
        artifact_id=artifact.id,
        version=number,
        storage_key=key,
        sha256=sha256_file(path),
        size=path.stat().st_size,
        job_id=job.id,
        qa_status=qa["status"],
        qa_report=qa,
        provenance=provenance,
    )
    db.add(version)
    db.flush()
    artifact.current_version_id = version.id
    audit(
        db,
        db.get(User, job.requested_by),
        "ARTIFACT_GENERATE",
        artifact.id,
        artifact.project_id,
        detail={"version": number},
    )
    return {
        "artifact_id": artifact.id,
        "version_id": version.id,
        "qa_status": qa["status"],
        "official_eligible": False,
    }


def qa_step(db, job, plan):
    from nps.pipeline import result_for

    result = result_for(db, job, "PPT")
    qa = qa_ppt(storage.path(result["storage_key"]), plan.data, template_for(db, plan).config)
    return persist_artifact(db, job, plan, "PPTX", result["storage_key"], qa, {})
