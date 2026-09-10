"""Editable editorial layouts and matching CPU video storyboard scenes."""

import json
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from nps.rendering import text_fits, font_file, wrap_text
from nps.errors import DomainError


def ink(color):
    rgb = [int(color[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    light = sum(
        w * (c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4)
        for w, c in zip((0.2126, 0.7152, 0.0722), rgb)
    )
    return "142D43" if light > 0.42 else "FFFFFF"


def mix(a, b, amount=0.1):
    return "".join(
        f"{round(int(a[i : i + 2], 16) * (1 - amount) + int(b[i : i + 2], 16) * amount):02X}"
        for i in (0, 2, 4)
    )


def word_lines(value, size, width, balanced=False):
    """Explicit word boundaries prevent Korean word fragments on the next line."""
    metric = ImageFont.truetype(font_file(), round(size * 96 / 72))
    available = (width - 0.12) * 96 * 0.90
    if balanced and "\n" not in value and metric.getlength(value) > available:
        words = value.split()
        candidates = [(" ".join(words[:i]), " ".join(words[i:])) for i in range(1, len(words))]
        fitting = [
            (a, b) for a, b in candidates if max(metric.getlength(a), metric.getlength(b)) <= available
        ]
        if fitting:
            return "\n".join(
                min(fitting, key=lambda pair: abs(metric.getlength(pair[0]) - metric.getlength(pair[1])))
            )
    output = []
    for paragraph in value.splitlines():
        line = ""
        for word in paragraph.split():
            candidate = (line + " " + word).strip()
            if line and metric.getlength(candidate) > available:
                output.append(line)
                line = word
            else:
                line = candidate
        output.append(line)
    return "\n".join(output)


def scene(item, policy, index=1, total=1):
    bg, fg, accent = (policy[k] for k in ("background", "foreground", "accent"))
    shapes = []

    def rect(x, y, w, h, color):
        shapes.append(dict(kind="rect", x=x, y=y, w=w, h=h, color=color))

    def text(value, x, y, w, h, size=24, color=None, bold=False, meta=False):
        minimum = 11 if meta else 18
        original = value
        value = word_lines(original, size, w, bold)
        while size > minimum and not text_fits(value, size, w - 0.12, h - 0.08):
            size -= 1
            value = word_lines(original, size, w, bold)
        shapes.append(
            dict(
                kind="text",
                text=value,
                x=x,
                y=y,
                w=w,
                h=h,
                size=size,
                color=color or fg,
                bold=bold,
                meta=meta,
            )
        )

    layout = item["layout_type"]
    title_x = max(0.65, min(1.1, policy.get("title_x", 0.75)))
    lines = [
        line
        for b in item["content_blocks"]
        if b["type"] == "text"
        for line in b.get("text", "").splitlines()
        if line.strip()
    ]
    if layout == "key_points" and policy.get("layout_arrangement") == "columns" and 2 <= len(lines) <= 3:
        layout = "comparison"
    if layout == "cover":
        rect(9.65, 0.65, 3.0, 5.8, accent)
        text("DOCUMENT / PRESENTATION", 0.75, 0.9, 8.4, 0.4, 12, meta=True)
        text(item["title"], 0.75, 2.0, 8.5, 2.3, 40, bold=True)
        text("\n".join(lines), 0.75, 4.7, 8.2, 1.25, 22)
        text(f"{total:02d}", 10.05, 3.0, 2.0, 1.2, 56, ink(accent), True)
    else:
        rect(0.75, 0.68, 0.55, 0.06, accent)
        text(
            item["title"], title_x, 0.9, 12.55 - title_x, 1.15, min(34, policy.get("title_pt", 32)), bold=True
        )
        structured = next(
            (b for b in item["content_blocks"] if b["type"] in {"table", "chart", "image"}), None
        )
        if sum(b["type"] in {"table", "chart", "image"} for b in item["content_blocks"]) > 1:
            raise DomainError("PPT_SPLIT_STRUCTURED_BLOCKS_REQUIRED", 422)
        if structured:
            shapes.append(
                dict(
                    kind=structured["type"],
                    x=0.8,
                    y=2.3,
                    w=11.7,
                    h=3.1 if lines else 4.0,
                    cells=structured.get("cells", []),
                )
            )
            if lines:
                text("\n".join(lines), 0.8, 5.55, 11.7, 0.95, 22)
        elif layout in {"comparison", "process"} and len(lines) >= 2:
            count = min(3, len(lines))
            gap = 0.45
            width = (11.7 - gap * (count - 1)) / count
            for i in range(count):
                x = 0.8 + i * (width + gap)
                rect(x, 2.45, width, 0.055, accent)
                text(f"{i + 1:02d}", x, 2.7, width, 0.6, 24, accent, True)
                if layout == "process" and i < count - 1:
                    text("→", x + width + 0.04, 2.8, 0.36, 0.5, 18, accent)
                # All remaining points stay visible in the last column.
                text("\n".join(lines[i:]) if i == count - 1 else lines[i], x, 3.5, width, 2.6, 24)
        else:
            count = max(1, len(lines))
            available = 4.25 / count
            for i, line in enumerate(lines):
                y = 2.2 + i * available
                text(f"{i + 1:02d}", 0.8, y, 0.65, 0.6, 23, accent, True)
                text(line, 1.8, y, 10.7, available - 0.12, policy.get("body_pt", 24))
                if i < count - 1:
                    rect(1.8, y + available - 0.10, 10.65, 0.012, mix(bg, fg, 0.16))
    text("발표용 핵심 발췌 · 원문 근거는 슬라이드 노트", 0.75, 6.65, 9.8, 0.32, 11, meta=True)
    text(f"{index:02d} / {total:02d}", 11.55, 6.65, 1.05, 0.32, 11, meta=True)
    return shapes


def font(paragraph, size, color, name, bold=False):
    paragraph.font.name = name
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    paragraph.line_spacing = 1.22
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    # Explicit East Asian typeface and neutral character spacing avoid theme inheritance.
    prop = paragraph._p.get_or_add_pPr().get_or_add_defRPr()
    prop.set("spc", "0")
    for tag in ("a:ea", "a:latin"):
        element = prop.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1])
        if element is None:
            element = OxmlElement(tag)
            prop.append(element)
        element.set("typeface", name)


def render(plan, output, policy, images=None):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    for number, item in enumerate(plan["slides"], 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(policy["background"])
        for part in scene(item, policy, number, len(plan["slides"])):
            x, y, w, h = [Inches(part[k]) for k in ("x", "y", "w", "h")]
            kind = part["kind"]
            if kind == "rect":
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
                shape.name = "decor:rule"
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(part["color"])
                shape.line.fill.background()
                shape._element.spPr.append(OxmlElement("a:effectLst"))
            elif kind == "text":
                shape = slide.shapes.add_textbox(x, y, w, h)
                shape.name = "meta:footer" if part["meta"] else "content:text"
                frame = shape.text_frame
                frame.word_wrap = True
                frame.margin_left = frame.margin_right = Inches(0.06)
                frame.margin_top = frame.margin_bottom = Inches(0.04)
                frame.text = part["text"]
                for p in frame.paragraphs:
                    font(p, part["size"], part["color"], policy["font"], part["bold"])
            elif kind == "image":
                path = (images or {}).get(item["slide_id"])
                if path:
                    with Image.open(path) as im:
                        ratio = im.width / im.height
                    width = min(part["w"], part["h"] * ratio)
                    height = width / ratio
                    slide.shapes.add_picture(
                        str(path),
                        Inches(part["x"] + (part["w"] - width) / 2),
                        y,
                        width=Inches(width),
                        height=Inches(height),
                    )
            elif kind == "table":
                cells = part["cells"]
                table = slide.shapes.add_table(len(cells), len(cells[0]), x, y, w, h).table
                for ri, row in enumerate(cells):
                    for ci, value in enumerate(row):
                        cell = table.cell(ri, ci)
                        cell.text = value
                        cell.margin_left = cell.margin_right = Inches(0.1)
                        cell.margin_top = cell.margin_bottom = Inches(0.04)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(
                            policy["accent"]
                            if ri == 0
                            else mix(policy["background"], policy["foreground"], 0.04 if ri % 2 else 0.09)
                        )
                        for p in cell.text_frame.paragraphs:
                            font(
                                p,
                                18,
                                ink(policy["accent"]) if ri == 0 else policy["foreground"],
                                policy["font"],
                                ri == 0,
                            )
            elif kind == "chart":
                cells = part["cells"]
                data = CategoryChartData()
                data.categories = [r[0] for r in cells[1:]]
                for col in range(1, len(cells[0])):
                    data.add_series(cells[0][col], [float(r[col]) for r in cells[1:]])
                chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, data).chart
                chart.font.name = policy["font"]
                chart.font.size = Pt(18)
                chart.font.color.rgb = RGBColor.from_string(policy["foreground"])
                chart.has_legend = len(cells[0]) > 2
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.chart_style = 10
                for si, series in enumerate(chart.series):
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = RGBColor.from_string(
                        mix(policy["accent"], policy["foreground"], si * 0.3)
                    )
        slide.notes_slide.notes_text_frame.text = json.dumps(
            dict(
                slide_id=item["slide_id"],
                source_refs=item["source_refs"],
                mock=plan["mock"],
                provenance=plan["provenance"],
            ),
            ensure_ascii=False,
        )
    deck.core_properties.title = plan["title"]
    deck.core_properties.subject = "Editorial prototype; source evidence in notes; official specification TBD"
    deck.save(output)


def preview(item, path, policy, image_path=None):
    """Same geometry/text as PPT; native chart is approximated with CPU bars."""
    canvas = Image.new("RGB", (1920, 1080), "#" + policy["background"])
    draw = ImageDraw.Draw(canvas)
    scale = 144
    for p in scene(item, policy, item.get("order", 1), policy.get("slide_count", 1)):
        x, y, w, h = [p[k] * scale for k in ("x", "y", "w", "h")]
        if p["kind"] == "text" and not text_fits(p["text"], p["size"], p["w"] - 0.12, p["h"] - 0.08):
            raise DomainError("VIDEO_CONTENT_OVERFLOW", 422)
        if p["kind"] == "image" and not image_path:
            raise DomainError("VIDEO_SOURCE_IMAGE_MISSING", 422)
        if p["kind"] == "rect":
            draw.rectangle((x, y, x + w, y + h), fill="#" + p["color"])
        elif p["kind"] == "text":
            f = ImageFont.truetype(font_file(), round(p["size"] * 2))
            for ri, line in enumerate(wrap_text(p["text"], f, w - 0.12 * scale)):
                draw.text(
                    (x + 0.06 * scale, y + 0.04 * scale + ri * p["size"] * 2 * 1.22),
                    line,
                    font=f,
                    fill="#" + p["color"],
                )
        elif p["kind"] == "image" and image_path:
            with Image.open(image_path) as im:
                im = im.convert("RGB")
                im.thumbnail((int(w), int(h)))
                canvas.paste(im, (round(x + (w - im.width) / 2), round(y)))
        elif p["kind"] == "table":
            cells = p["cells"]
            cw = w / len(cells[0])
            rh = h / len(cells)
            f = ImageFont.truetype(font_file(), 36)
            for ri, row in enumerate(cells):
                for ci, value in enumerate(row):
                    color = (
                        policy["accent"]
                        if ri == 0
                        else mix(policy["background"], policy["foreground"], 0.04 if ri % 2 else 0.09)
                    )
                    draw.rectangle(
                        (x + ci * cw, y + ri * rh, x + (ci + 1) * cw, y + (ri + 1) * rh), fill="#" + color
                    )
                    for li, line in enumerate(wrap_text(value, f, cw - 28)):
                        draw.text(
                            (x + ci * cw + 14, y + ri * rh + 6 + li * 44),
                            line,
                            font=f,
                            fill="#" + (ink(color) if ri == 0 else policy["foreground"]),
                        )
        elif p["kind"] == "chart":
            cells = p["cells"]
            series = len(cells[0]) - 1
            values = [float(v) for r in cells[1:] for v in r[1:]]
            low = min(0, min(values))
            high = max(1, max(values))
            span = high - low
            plot_h = h - 120
            baseline = y + 30 + high / span * plot_h
            cw = w / (len(cells) - 1)
            bw = (cw - 32) / series
            f = ImageFont.truetype(font_file(), 26)
            draw.line((x, baseline, x + w, baseline), fill="#" + policy["foreground"], width=2)
            for i, row in enumerate(cells[1:]):
                for si, value in enumerate(row[1:]):
                    v = float(value)
                    top = y + 30 + (high - v) / span * plot_h
                    left = x + i * cw + 16 + si * bw
                    color = mix(policy["accent"], policy["foreground"], si * 0.3)
                    draw.rectangle(
                        (left, min(top, baseline), left + bw - 8, max(top, baseline) + 1), fill="#" + color
                    )
                    draw.text((left, min(top, baseline) - 30), value, font=f, fill="#" + policy["foreground"])
                draw.text((x + i * cw + 16, y + h - 55), row[0], font=f, fill="#" + policy["foreground"])
            if series > 1:
                draw.text((x, y - 35), " / ".join(cells[0][1:]), font=f, fill="#" + policy["foreground"])
    canvas.save(path, "PNG")
