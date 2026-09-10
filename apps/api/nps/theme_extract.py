"""Extract bounded style statistics only; never execute or copy example slide content."""

import json
import os
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path
from nps.config import settings
from nps.errors import DomainError
from nps.storage import storage
from nps.themes import BASE


def extract(path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor

    deck = Presentation(path)
    if not 1 <= len(deck.slides) <= 200:
        raise DomainError("THEME_SLIDE_LIMIT", 422)
    colors, fonts, sizes, layouts, backgrounds = Counter(), Counter(), [], [], Counter()
    # Most real PPTs inherit colors from the master instead of writing RGB per run.
    from defusedxml.ElementTree import fromstring

    theme_colors = []
    theme_background = None
    for rel in deck.slide_master.part.rels.values():
        if rel.reltype.endswith("/theme") and not rel.is_external:
            root = fromstring(rel.target_part.blob)
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            scheme = root.find(".//a:clrScheme", ns)
            if scheme is not None:
                for slot in scheme:
                    if not len(slot):
                        continue
                    value = slot[0].get("lastClr") or slot[0].get("val", "")
                    if re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                        value = value.upper()
                        if slot.tag.endswith("}lt1"):
                            theme_background = value
                        if "accent" in slot.tag:
                            theme_colors.append(value)

    def color(value):
        try:
            return str(value.rgb) if isinstance(value.rgb, RGBColor) else None
        except (AttributeError, TypeError):
            return None

    shape_count = 0
    for slide in deck.slides:
        try:
            bg = color(slide.background.fill.fore_color)
            if bg:
                backgrounds[bg] += 1
        except (TypeError, AttributeError):
            pass
        pattern = []
        for shape in slide.shapes:
            shape_count += 1
            if shape_count > 15000:
                raise DomainError("THEME_SHAPE_LIMIT", 422)
            kind = (
                "text"
                if shape.has_text_frame
                else "table"
                if shape.has_table
                else "image"
                if shape.shape_type == 13
                else "shape"
            )
            pattern.append(
                {
                    "kind": kind,
                    "x": round(shape.left / deck.slide_width, 3),
                    "y": round(shape.top / deck.slide_height, 3),
                    "w": round(shape.width / deck.slide_width, 3),
                    "h": round(shape.height / deck.slide_height, 3),
                }
            )
            try:
                value = color(shape.fill.fore_color)
                if value:
                    colors[value] += 1
            except (TypeError, AttributeError):
                pass
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for font in [paragraph.font, *[r.font for r in paragraph.runs]]:
                        if font.name:
                            fonts[font.name[:80]] += 1
                        if font.size and 10 <= font.size.pt <= 80:
                            sizes.append(font.size.pt)
                        value = color(font.color)
                        if value:
                            colors[value] += 1
        layouts.append(pattern[:80])
    # Choose chromatic accents, keeping a readable background/foreground pair.
    palette = list(
        dict.fromkeys(
            [c for c, _ in colors.most_common(16) if re.fullmatch(r"[0-9A-F]{6}", c)] + theme_colors
        )
    )[:20]
    bg = backgrounds.most_common(1)[0][0] if backgrounds else theme_background or "F8FAFC"
    rgb = [int(bg[i : i + 2], 16) for i in (0, 2, 4)]
    fg = "F4F7FA" if sum(rgb) < 360 else "142D43"
    accents = [
        c
        for c in palette
        if max(int(c[i : i + 2], 16) for i in (0, 2, 4)) - min(int(c[i : i + 2], 16) for i in (0, 2, 4)) > 45
        and c != bg
    ]
    title_size = max(30, min(38, sorted(sizes)[int(len(sizes) * 0.8)])) if sizes else 32
    text_patterns = [
        p for layout in layouts for p in layout if p["kind"] == "text" and p["y"] < 0.22 and p["w"] > 0.3
    ]
    title_x = max(0.65, min(1.1, min((p["x"] for p in text_patterns), default=0.05) * 13.333))
    columns = any(
        sum(p["kind"] == "text" and p["y"] > 0.25 and 0.15 < p["w"] < 0.45 for p in layout) >= 2
        for layout in layouts
    )
    return {
        **BASE,
        "layout_arrangement": "columns" if columns else "rows",
        "background": bg,
        "foreground": fg,
        "accent": accents[0] if accents else "257C88",
        "title_pt": title_size,
        "title_x": round(title_x, 2),
        "palette": palette,
        "reference_fonts": dict(fonts.most_common(6)),
        "reference_layouts": layouts[:40],
        "layout_preferences": [
            "image" if sum(p["kind"] == "image" for p in layout) else "key_points" for layout in layouts[:12]
        ],
        "reference_slides": len(deck.slides),
        "learning_version": "style-features-2.0",
        "source": "uploaded-reference",
        "warnings": [
            "글꼴은 설치된 한국어 폰트로 정규화합니다.",
            "원본 본문·로고·사진을 새 발표자료에 복사하지 않습니다.",
        ],
        "learning_method": "persistent-style-retrieval; not model fine-tuning",
    }


def safe_extract(path):
    _, output = storage.allocate("parsed", ".json")
    env = os.environ.copy()
    env["PYTHONPATH"] = str(Path(__file__).resolve().parents[1])
    try:
        result = subprocess.run(
            [sys.executable, "-m", "nps.theme_extract", str(path.resolve()), str(output.resolve())],
            capture_output=True,
            env=env,
            timeout=settings().parse_timeout + 5,
        )
        data = json.loads(output.read_text(encoding="utf-8")) if output.exists() else {}
        if result.returncode or "background" not in data:
            raise DomainError(data.get("error", "THEME_PARSE_FAILED"), 422)
        return data
    except subprocess.TimeoutExpired:
        raise DomainError("THEME_PARSE_TIMEOUT", 422) from None
    finally:
        output.unlink(missing_ok=True)


if __name__ == "__main__":
    if os.name == "posix":
        import resource

        limit = settings().parse_memory_mb * 1024 * 1024
        resource.setrlimit(resource.RLIMIT_AS, (limit, limit))
        resource.setrlimit(resource.RLIMIT_CPU, (settings().parse_timeout, settings().parse_timeout + 1))
    try:
        data, code = extract(Path(sys.argv[1])), 0
    except Exception as exc:
        data, code = {"error": exc.code if isinstance(exc, DomainError) else "THEME_PARSE_FAILED"}, 1
    Path(sys.argv[2]).write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    sys.exit(code)
